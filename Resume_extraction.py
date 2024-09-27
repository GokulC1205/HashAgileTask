import os
import streamlit as st
from pdfminer.high_level import extract_text
import docx
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.messages import HumanMessage
from texttable import Texttable
import re
from dotenv import load_dotenv
import pandas as pd
from fpdf import FPDF
from pymongo import MongoClient


load_dotenv()

mongo_uri = "mongodb://localhost:27017"  
client = MongoClient(mongo_uri)
db = client['resume_db'] 
collection = db['resumes']  

gemini_api_key = os.getenv("API_KEY")  

def extract_text_from_pdf(pdf_path):
    return extract_text(pdf_path)

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    full_text = [para.text for para in doc.paragraphs]
    return '\n'.join(full_text)

def extract_text_from_txt(txt_path):
    with open(txt_path, 'r', encoding='utf-8') as file:
        return file.read()

def resume_to_text(file):
    file_extension = os.path.splitext(file.name)[1].lower()
    
    if file_extension == '.pdf':
        resume_text = extract_text_from_pdf(file)
    elif file_extension == '.docx':
        resume_text = extract_text_from_docx(file)
    elif file_extension == '.txt':
        resume_text = extract_text_from_txt(file)
    else:
        st.error("Unsupported file format. Please use PDF, DOCX, or TXT.")
        return None
    return resume_text

st.title('Resume Information Extractor')

uploaded_file = st.file_uploader("Upload a resume (PDF, DOCX, or TXT)", type=['pdf', 'docx', 'txt'])

if uploaded_file:
    with st.spinner('Extracting text from the resume...'):
        resume_text = resume_to_text(uploaded_file)
    
    if resume_text:
        if gemini_api_key:
            llm = ChatGoogleGenerativeAI(api_key=gemini_api_key, model="gemini-pro")

            prompt_message = HumanMessage(
                content=(
                    f"Extract the following information from this resume: name, email, phone number, college name, and skills. "
                    f"Please return the information in the following format:\n"
                    f"Name: [name]\nEmail: [email]\nPhone Number: [phone number]\nCollege Name: [college name]\nSkills: [skills]\n\n"
                    f"Here is the resume text:\n{resume_text}"
                )
            )

            with st.spinner('Extracting information from the resume...'):
                output = llm.invoke([prompt_message])
            
            extracted_content = output.content
            name, email, phone_number, college_name, skills = None, None, None, None, None
            for line in extracted_content.split('\n'):
                if "Name:" in line and not name:
                    name = line.split("Name:")[-1].lstrip('*').strip()
                elif "Email:" in line and not email:
                    email = re.findall(r'\S+@\S+', line)
                    email = email[0] if email else None
                elif "Phone Number:" in line and not phone_number:
                    phone_number = re.findall(r'\+?\d[\d -]{8,12}\d', line)
                    phone_number = phone_number[0] if phone_number else None
                elif "College Name:" in line and not college_name:
                    college_name = line.split("College Name:")[-1].lstrip('*').strip()
                elif "Skills:" in line:
                    skills_section = extracted_content.split("Skills:")[-1].lstrip('*').strip()
                    skills = '\n'.join([s.lstrip('-').strip() for s in skills_section.split('\n') if s.strip()])
            
            resume_data = {
                "name": name,
                "email": email,
                "phone_number": phone_number,
                "college_name": college_name,
                "skills": skills
            }

            collection.insert_one(resume_data)

            st.success("Resume data inserted into MongoDB successfully!")

            df = pd.DataFrame({
                "Field": ["Name", "Email", "Phone Number", "College Name", "Skills"],
                "Information": [name, email, phone_number, college_name, skills]
            })
            table = Texttable()
            table.set_cols_width([20, 100])
            table.set_cols_dtype(["t", "t"])
            table.add_row(["Field", "Information"])
            table.add_row(["Name", name])
            table.add_row(["Email", email])
            table.add_row(["Phone Number", phone_number])
            table.add_row(["College Name", college_name])
            table.add_row(["Skills", skills])

            st.subheader('Extracted Information')
            st.text(table.draw())

            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            pdf.cell(200, 10, txt="Extracted Resume Information", ln=True, align='C')
            pdf.ln(10)
            pdf.cell(40, 10, f"Name: {name}", ln=True)
            pdf.cell(40, 10, f"Email: {email}", ln=True)
            pdf.cell(40, 10, f"Phone Number: {phone_number}", ln=True)
            pdf.cell(40, 10, f"College Name: {college_name}", ln=True)
            pdf.cell(40, 10, "Skills:", ln=True)
            pdf.multi_cell(0, 10, skills)

            pdf_file = "resume_info.pdf"
            pdf.output(pdf_file, 'F')

            with open(pdf_file, "rb") as f:
                pdf_data = f.read()

            st.download_button(
                label="Download Resume Info as PDF",
                data=pdf_data,
                file_name=pdf_file,
                mime="application/pdf",
            )
        else:
            st.error("Gemini API key not found. Please check your environment settings.")

else:
    st.info("Please upload a resume to extract information.")
